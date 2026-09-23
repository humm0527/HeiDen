"""将股票打分 Excel 的月截面展开为按季度保存的逐交易日集中度 CSV。

首次生成：2026-09-17。日期边界采用上一截面包含下一截面日期的规则。
依赖：Python >= 3.10、openpyxl。用法见同目录 README.md。
"""

from __future__ import annotations

import argparse
from bisect import bisect_left, bisect_right
import csv
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import sys
from tempfile import TemporaryDirectory

import openpyxl


SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR.parent
OUTPUT_COLUMNS = ["biz_date", "stk_code", "证券名称", "券商03"]
INPUT_COLUMNS = ["证券代码", "证券名称", "集中度分类"]
FILE_PATTERN = re.compile(r"股票打分结果(\d{8})\.xlsx$")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_calendar(path: Path, market: str = "XSHG") -> list[date]:
    """读取显式交易日列表，不用工作日近似替代交易日。"""
    days = set()
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not {"market_code", "calendar_date", "is_trading_day"} <= set(reader.fieldnames or []):
            raise ValueError("交易日历必须包含 market_code,calendar_date,is_trading_day")
        for row in reader:
            if row["market_code"] != market:
                continue
            flag = row["is_trading_day"].strip().lower()
            if flag not in {"true", "false", "1", "0"}:
                raise ValueError(f"无效交易日标志：{row}")
            if flag in {"true", "1"}:
                days.add(date.fromisoformat(row["calendar_date"]))
    if not days:
        raise ValueError(f"交易日历中没有 {market} 的交易日")
    return sorted(days)


def previous_or_same_day(value: date, calendar: list[date]) -> date:
    # 拒绝把超出日历范围的日期悄悄回退到最后一天。
    if not calendar[0] <= value <= calendar[-1]:
        raise ValueError(f"日期 {value} 超出交易日历范围 {calendar[0]}—{calendar[-1]}，请更新日历")
    return calendar[bisect_right(calendar, value) - 1]


def discover_snapshots(input_dir: Path, calendar: list[date]) -> list[dict]:
    snapshots = []
    for path in sorted(input_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        match = FILE_PATTERN.fullmatch(path.name)
        if not match:
            raise ValueError(f"输入文件名应为 股票打分结果YYYYMMDD.xlsx：{path.name}")
        raw_day = datetime.strptime(match[1], "%Y%m%d").date()
        snapshots.append({"path": path, "file_date": raw_day,
                          "adjusted_date": previous_or_same_day(raw_day, calendar)})
    if not snapshots:
        raise ValueError(f"未找到输入文件：{input_dir}")
    boundaries = [s["adjusted_date"] for s in snapshots]
    if len(set(boundaries)) != len(boundaries):
        raise ValueError("多个截面回溯到了同一交易日，无法确定优先级")
    return snapshots


def source_index(day: date, boundaries: list[date]) -> int:
    """边界当天仍用上一截面；首个截面从其边界当天开始使用。"""
    if not boundaries or day < boundaries[0]:
        raise ValueError(f"{day} 早于首个截面，不能向历史回填未来数据")
    return max(0, bisect_left(boundaries, day) - 1)


def read_snapshot(path: Path, blank_policy: str, sheet_name: str = "Sheet1") -> tuple[list, int]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{path.name} 缺少工作表 {sheet_name}")
        source = workbook[sheet_name].iter_rows(values_only=True)
        header = list(next(source, ()))
        if any(header.count(name) != 1 for name in INPUT_COLUMNS):
            raise ValueError(f"{path.name} 必须各有一列：{INPUT_COLUMNS}")
        indices = [header.index(name) for name in INPUT_COLUMNS]
        result, seen, blank_count = [], set(), 0
        for line, row in enumerate(source, 2):
            if all(value is None for value in row):
                continue
            code, name, category = [row[i] for i in indices]
            if not isinstance(code, str) or not re.fullmatch(r"\d{6}\.(SH|SZ|BJ)", code):
                raise ValueError(f"{path.name} 第 {line} 行证券代码无效：{code!r}")
            if code in seen:
                raise ValueError(f"{path.name} 重复证券代码：{code}")
            seen.add(code)
            if not isinstance(name, str) or not name.strip():
                raise ValueError(f"{path.name} 第 {line} 行证券名称为空")
            if category is None or (isinstance(category, str) and not category.strip()):
                blank_count += 1
                if blank_policy == "error":
                    raise ValueError(f"{path.name} 第 {line} 行分类为空，请指定 --blank-policy keep/g/drop")
                if blank_policy == "drop":
                    continue
                category = "G" if blank_policy == "g" else ""
            elif category not in set("ABCDEFG"):
                raise ValueError(f"{path.name} 第 {line} 行分类无效：{category!r}")
            result.append((code, name, category))
        if not result:
            raise ValueError(f"{path.name} 没有可输出的股票")
        return result, blank_count
    finally:
        workbook.close()


def convert(input_dir: Path, output_dir: Path, calendar_path: Path,
            blank_policy: str, start: date | None = None, end: date | None = None,
            market: str = "XSHG", sheet_name: str = "Sheet1") -> dict:
    calendar = load_calendar(calendar_path, market)
    snapshots = discover_snapshots(input_dir, calendar)
    boundaries = [s["adjusted_date"] for s in snapshots]
    start = start or boundaries[0]
    requested_end = end or snapshots[-1]["file_date"]
    end = previous_or_same_day(requested_end, calendar)
    if start < boundaries[0] or start > end:
        raise ValueError("开始日不能早于首个截面，也不能晚于结束日")
    days = [d for d in calendar if start <= d <= end]
    if not days:
        raise ValueError("所选区间没有交易日")
    if output_dir.resolve() in {input_dir.resolve(), (DATA_DIR / "history_single").resolve()}:
        raise ValueError("输出目录不能覆盖源目录或参考模板目录")

    # 在写出任何正式文件前，先验证所有源文件（包含尚未生效的最后一份）。
    records = []
    for snapshot in snapshots:
        rows, blank_count = read_snapshot(snapshot["path"], blank_policy, sheet_name)
        records.append(rows)
        snapshot.update(stock_count=len(rows), blank_class_count=blank_count,
                        sha256=sha256(snapshot["path"]), output_days=[])

    warnings = []
    for previous, current in zip(snapshots, snapshots[1:]):
        a, b = previous["file_date"], current["file_date"]
        month_gap = (b.year - a.year) * 12 + b.month - a.month
        if month_gap > 1:
            warnings.append(f"{a} 至 {b} 之间存在缺失月份，按相邻文件规则延续前一截面")

    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []
    # 临时目录中完整写出，成功后替换本次涉及的季度文件。
    with TemporaryDirectory(prefix=".monthly_to_daily_", dir=output_dir) as temporary:
        quarters = sorted({(d.year, (d.month - 1) // 3 + 1) for d in days})
        for year, quarter in quarters:
            quarter_days = [d for d in days if (d.year, (d.month - 1) // 3 + 1) == (year, quarter)]
            filename = f"集中度分类_{year}Q{quarter}.csv"
            temporary_path = Path(temporary) / filename
            row_count = 0
            with temporary_path.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle, lineterminator="\r\n")
                writer.writerow(OUTPUT_COLUMNS)
                for day in quarter_days:
                    index = source_index(day, boundaries)
                    writer.writerows((day.isoformat(), *row) for row in records[index])
                    row_count += len(records[index])
                    snapshots[index]["output_days"].append(day.isoformat())
            outputs.append({"file": filename, "first_date": quarter_days[0].isoformat(),
                            "last_date": quarter_days[-1].isoformat(),
                            "trading_days": len(quarter_days), "rows": row_count,
                            "sha256": sha256(temporary_path)})
        for output in outputs:
            (Path(temporary) / output["file"]).replace(output_dir / output["file"])

    source_manifest = []
    for snapshot in snapshots:
        used_days = snapshot["output_days"]
        source_manifest.append({
            "file": snapshot["path"].name,
            "file_date": snapshot["file_date"].isoformat(),
            "adjusted_date": snapshot["adjusted_date"].isoformat(),
            "effective_start": used_days[0] if used_days else None,
            "effective_end": used_days[-1] if used_days else None,
            "trading_days": len(used_days), "stock_count": snapshot["stock_count"],
            "blank_class_count": snapshot["blank_class_count"], "sha256": snapshot["sha256"],
        })
    return {"input_dir": str(input_dir.resolve()), "output_dir": str(output_dir.resolve()),
            "calendar": str(calendar_path.resolve()), "calendar_sha256": sha256(calendar_path),
            "calendar_market": market, "calendar_first": calendar[0].isoformat(),
            "calendar_last": calendar[-1].isoformat(), "blank_policy": blank_policy,
            "requested_end": requested_end.isoformat(), "first_date": days[0].isoformat(),
            "last_date": days[-1].isoformat(), "trading_days": len(days),
            "rows": sum(o["rows"] for o in outputs), "warnings": warnings,
            "sources": source_manifest, "outputs": outputs}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DATA_DIR / "股票打分结果")
    parser.add_argument("--output-dir", type=Path, default=DATA_DIR / "single")
    parser.add_argument("--calendar", type=Path, default=SCRIPT_DIR / "trading_calendar.csv")
    parser.add_argument("--market", default="XSHG")
    parser.add_argument("--sheet", default="Sheet1")
    parser.add_argument("--blank-policy", choices=["keep", "g", "drop", "error"], default="g")
    parser.add_argument("--start-date", type=date.fromisoformat)
    parser.add_argument("--end-date", type=date.fromisoformat)
    parser.add_argument("--manifest", type=Path, default=SCRIPT_DIR / "last_run_manifest.json")
    args = parser.parse_args()
    try:
        result = convert(args.input_dir, args.output_dir, args.calendar, args.blank_policy,
                         args.start_date, args.end_date, args.market, args.sheet)
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (ValueError, OSError) as exc:
        parser.exit(1, f"转换失败：{exc}\n")
    for warning in result["warnings"]:
        print(f"提示：{warning}", file=sys.stderr)
    for output in result["outputs"]:
        print(f"{output['file']}：{output['trading_days']} 个交易日，{output['rows']:,} 行")
    print(f"完成：{result['first_date']}—{result['last_date']}，"
          f"共 {result['trading_days']} 个交易日，{result['rows']:,} 行。")
    print(f"转换明细：{args.manifest}")


if __name__ == "__main__":
    main()
